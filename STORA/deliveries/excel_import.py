"""Excel import for Deliveries -- lets a warehouse clerk load a whole batch
of goods from a spreadsheet instead of adding products to a delivery one by
one. Parsing/matching only lives here; nothing in this module touches the
database -- it just builds what the delivery grid should show. Actual
product creation happens at Save time (see deliveries.views.deliveries_add),
so reviewing/editing/cancelling the import never leaves anything behind.

Column layout (see build_template_workbook) -- only Name is required, every
other column may be blank:
  Code | Name | Unit Type | Quantity | Category | Subcategory |
  Delivery Price | Sell Price | Main Supplier | Barcode 1 | Barcode 2 |
  Barcode 3

Unit Type ("pcs"/"kg", Bulgarian "бр"/"кг" also accepted, case-insensitive)
only matters for a genuinely NEW product -- it's ignored for a code/barcode
match, since an existing product's unit_type is already fixed and this file
is never allowed to edit the catalog (see matching rules below). Left blank
or unrecognised, a new product defaults to "pcs" -- this used to be silently
hardcoded to "pcs" for every imported product regardless of what it
actually was, which is wrong for anything sold by weight (produce, deli):
Product.quantity's step (whole units vs fractional kg) and scale-barcode
decoding (products/scale_barcode.py) both depend on unit_type being right
from the moment the product is created, not fixed up later by hand.

Matching rules (worked out with the shop owner before building this):
- Code or any of the 3 barcodes matching an existing product means this row
  IS that product -- a positive quantity delivers more of it, a zero
  quantity is a no-op (the product already exists, nothing to do), never a
  price/category/etc. edit of the existing product (the file is for
  receiving goods, not bulk-editing the catalog). Exact code/barcode
  matches never go through the fuzzy layers below at all.
- Category/Subcategory/Main Supplier are matched by name via Postgres
  trigram similarity (same tech as every other fuzzy picker in this app) --
  never auto-created when nothing matches confidently enough, just left
  blank. Getting a category wrong by guessing would make a worse mess than
  leaving it for a human to fill in.
- A row with no code/barcode match is a candidate NEW product -- also gets
  a fuzzy name check against the whole catalog (including archived
  products, since "this barcode changed but it's the same real item,
  possibly discontinued/renamed" is exactly the case worth catching) to
  flag a likely duplicate. Never auto-links either -- only flags it for the
  person reviewing the delivery grid to confirm or dismiss.

Two-layer matching (the original design discussed with the shop owner,
built in two parts):
- Layer 1 (always on, no external dependency): pure trigram similarity
  against two thresholds. >= SIMILARITY_HIGH is confident enough to use
  outright. < SIMILARITY_LOW isn't even a candidate. On its own, Layer 1
  produces a complete, working import with nothing in the middle band ever
  accepted -- conservative by design (see "Getting a category wrong..."
  above), the same philosophy just applied via a wider reject zone instead
  of a single cutoff.
- Layer 2 (only if ANTHROPIC_API_KEY is configured): for the middle band
  (similarity neither high enough for certainty nor low enough to dismiss
  outright), let Claude judge the specific pair more finely -- a genuinely
  additional hint on the borderline cases Layer 1 alone would conservatively
  drop, never the deciding factor on a case Layer 1 already resolved. Every
  ambiguous pair across the whole file is judged in ONE batched API call
  (see _ai_judge_ambiguous_matches), not one call per cell -- a file with
  many borderline rows would otherwise mean dozens of slow, costly round
  trips. If the call fails or no key is configured, every pending item
  simply falls back to Layer 1's own conservative default (reject) -- an
  AI hiccup degrades the import to "slightly more conservative", never
  breaks it.
"""

import json
from decimal import Decimal, InvalidOperation

import anthropic
import openpyxl
from django.conf import settings
from django.contrib.postgres.search import TrigramSimilarity

from STORA.products.models import Barcode, Category, Product, Suppliers

COLUMN_HEADERS = [
    'Code', 'Name', 'Unit Type', 'Quantity', 'Category', 'Subcategory',
    'Delivery Price', 'Sell Price', 'Main Supplier',
    'Barcode 1', 'Barcode 2', 'Barcode 3',
]

# Accepted spellings for the Unit Type column, both languages, lower-cased.
_WEIGHT_SPELLINGS = {'kg', 'weight', 'кг', 'тегло', 'теглов', 'теглови'}
_PIECE_SPELLINGS = {'pcs', 'pc', 'piece', 'бр', 'бройка', 'бройки', 'бройков', 'бройкови'}


def _clean_unit_type(value):
    """Excel's Unit Type cell -> Product.WEIGHT/Product.PIECE. Blank or
    unrecognised falls back to Product.PIECE (the same default Product
    itself uses), rather than erroring out -- most goods really are sold
    by the piece, and this column is optional like every other one."""
    text = _clean_str(value).lower()
    if text in _WEIGHT_SPELLINGS:
        return Product.WEIGHT
    return Product.PIECE


# Layer 1's two thresholds -- empirical, not derived from anything. Below
# SIMILARITY_LOW: not even a candidate, Layer 2 never sees it (asking Claude
# to weigh in on genuinely unrelated names would just waste a call and risk
# a false positive). At/above SIMILARITY_HIGH: confident enough on trigram
# alone, Layer 2 is skipped for this pair entirely -- it only ever runs for
# the band in between.
SIMILARITY_LOW = 0.30
SIMILARITY_HIGH = 0.55


def build_template_workbook():
    """Blank .xlsx with just the header row -- served by
    delivery_import_template for "Download template"."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Import'
    ws.append(COLUMN_HEADERS)
    for col_index in range(1, len(COLUMN_HEADERS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_index)].width = 16
    return wb


def _clean_str(value):
    if value is None:
        return ''
    return str(value).strip()


def _clean_decimal(value):
    if value is None or value == '':
        return None
    try:
        return Decimal(str(value).strip().replace(',', '.'))
    except InvalidOperation:
        return None


def _best_name_match(name, queryset):
    """Best row in `queryset` whose `name` field is trigram-similar to
    `name`, restricted to at least SIMILARITY_LOW -- below that it isn't a
    candidate at all, for either layer. Returns (candidate, similarity) or
    (None, None); `queryset` must be a model with a `name` field. Doesn't
    decide anything on its own -- see _classify_similarity."""
    if not name:
        return None, None
    match = (
        queryset.annotate(similarity=TrigramSimilarity('name', name))
        .filter(similarity__gte=SIMILARITY_LOW)
        .order_by('-similarity')
        .first()
    )
    if not match:
        return None, None
    return match, match.similarity


def _classify_similarity(similarity):
    """'confident' (use the match outright, Layer 2 not needed) /
    'ambiguous' (Layer 2's territory) / None (no candidate at all --
    _best_name_match already returns None below SIMILARITY_LOW)."""
    if similarity is None:
        return None
    return 'confident' if similarity >= SIMILARITY_HIGH else 'ambiguous'


def _ai_judge_ambiguous_matches(pending):
    """Layer 2. `pending`: list of {'key', 'kind', 'name', 'candidate_name'}
    dicts, one per ambiguous pair across the WHOLE file. Returns {key: bool}
    -- True means Claude judges `name` (from the Excel file) and
    `candidate_name` (already in the catalog) most likely refer to the same
    real-world product/category/supplier, just written differently. Missing
    keys (no API key, empty `pending`, a malformed/failed response) simply
    mean "no Layer 2 verdict for this one" -- the caller falls back to
    Layer 1's own conservative default, never raises."""
    if not pending:
        return {}
    api_key = settings.ANTHROPIC_API_KEY
    if not api_key:
        return {}

    lines = [
        f'{i}. ({item["kind"]}) Excel text: "{item["name"]}" -- existing catalog entry: "{item["candidate_name"]}"'
        for i, item in enumerate(pending)
    ]
    prompt = (
        'For each numbered pair below, decide whether the Excel text and the '
        'existing catalog entry most likely refer to the SAME real-world '
        'product, category, or supplier -- just written differently '
        '(reordered words, abbreviation, minor spelling difference, '
        'different case) -- or whether they are actually DIFFERENT things '
        'that just happen to share some letters. The shop is in Bulgaria; '
        'both Bulgarian and English text occur, sometimes mixed.\n\n'
        + '\n'.join(lines)
        + '\n\nRespond with ONLY a JSON array of true/false, exactly '
        f'{len(pending)} values, same order as the numbered list above -- '
        'e.g. [true, false, true]. No other text, no explanation.'
    )
    try:
        client = anthropic.Anthropic(api_key=api_key)
        response = client.messages.create(
            model=settings.AI_REPORTS_MODEL,
            max_tokens=1024,
            messages=[{'role': 'user', 'content': prompt}],
        )
        text = ''.join(block.text for block in response.content if block.type == 'text').strip()
        verdicts = json.loads(text)
    except Exception:
        # Never let an AI hiccup (bad key, network error, malformed JSON
        # back) break the import -- Layer 2 is a bonus on top of a already-
        # complete Layer 1, not a requirement. See module docstring.
        return {}

    if not isinstance(verdicts, list) or len(verdicts) != len(pending):
        return {}
    return {pending[i]['key']: bool(verdicts[i]) for i in range(len(pending))}


def parse_delivery_import(file_obj):
    """Reads an uploaded .xlsx and resolves every row against the current
    catalog. Returns {'rows': [...], 'skipped': [...], 'errors': [...]}.

    Each entry in `rows` is either:
      {'row_number', 'is_new': False, 'product_id', 'quantity'}
    for a code/barcode match, or:
      {'row_number', 'is_new': True, 'internal_code', 'name', 'unit_type',
       'quantity', 'delivery_price', 'sell_price', 'category_id',
       'category_name', 'supplier_id', 'supplier_name', 'barcodes': [...],
       'duplicate_candidate': {'id', 'name'} | None}
    for a genuinely new product candidate.

    Two passes: the first resolves every row's Layer-1 matches and collects
    whatever fell in the ambiguous middle band; ONE batched Layer-2 call (if
    configured) then judges every ambiguous pair across the whole file at
    once; the second pass finalises each row using that verdict (or Layer
    1's conservative default where there isn't one)."""
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    except Exception:
        return {'rows': [], 'skipped': [], 'errors': ['Could not read this file -- is it a real .xlsx?']}
    ws = wb.active

    pending_rows = []
    skipped = []
    errors = []
    ai_pending = []
    next_code_counter = None

    all_rows = list(ws.iter_rows(values_only=True))
    if len(all_rows) <= 1:
        return {'rows': [], 'skipped': [], 'errors': ['The file has no data rows.']}

    for row_number, raw_row in enumerate(all_rows[1:], start=2):
        if raw_row is None or all(cell in (None, '') for cell in raw_row):
            continue

        cells = list(raw_row) + [None] * (len(COLUMN_HEADERS) - len(raw_row))
        (code, name, unit_type, quantity, category_name, subcategory_name,
         delivery_price, sell_price, supplier_name, bc1, bc2, bc3) = cells[:12]

        code = _clean_str(code)
        name = _clean_str(name)
        if not name:
            errors.append(f'Row {row_number}: name is required, row skipped.')
            continue

        unit_type = _clean_unit_type(unit_type)
        quantity = _clean_decimal(quantity) or Decimal('0')
        delivery_price = _clean_decimal(delivery_price)
        sell_price = _clean_decimal(sell_price)
        barcodes = [b for b in (_clean_str(bc1), _clean_str(bc2), _clean_str(bc3)) if b]

        matched_product = None
        if code:
            matched_product = Product.objects.filter(internal_code=code).first()
        if not matched_product and barcodes:
            barcode_match = Barcode.objects.filter(code__in=barcodes).select_related('product').first()
            if barcode_match:
                matched_product = barcode_match.product

        if matched_product:
            if quantity <= 0:
                skipped.append(f'Row {row_number}: "{matched_product.name}" already exists, skipped (quantity was 0).')
                continue
            pending_rows.append({
                'row_number': row_number,
                'is_new': False,
                'product_id': matched_product.pk,
                'quantity': float(quantity),
            })
            continue

        duplicate_match, duplicate_sim = _best_name_match(name, Product.objects.all())

        category_match, category_sim = (
            _best_name_match(category_name, Category.objects.filter(parent__isnull=True))
            if category_name else (None, None)
        )
        # Scoped by the RAW category candidate regardless of how confident
        # it is -- narrowing which subcategories even get considered, not a
        # final decision (that happens in the second pass below), so it
        # mirrors the pre-Layer-2 behaviour exactly.
        subcategory_match, subcategory_sim = (None, None)
        if subcategory_name:
            sub_qs = Category.objects.filter(parent_id=category_match.pk) if category_match \
                else Category.objects.exclude(parent__isnull=True)
            subcategory_match, subcategory_sim = _best_name_match(subcategory_name, sub_qs)

        supplier_match, supplier_sim = (
            _best_name_match(supplier_name, Suppliers.objects.all()) if supplier_name else (None, None)
        )

        if not code:
            if next_code_counter is None:
                from STORA.products.views import get_free_internal_codes
                next_code_counter = get_free_internal_codes()[1]
            code = str(next_code_counter)
            next_code_counter += 1

        row = {
            'row_number': row_number,
            'is_new': True,
            'internal_code': code,
            'name': name,
            'unit_type': unit_type,
            'quantity': float(quantity),
            'delivery_price': float(delivery_price) if delivery_price is not None else None,
            'sell_price': float(sell_price) if sell_price is not None else None,
            'barcodes': barcodes,
            # Resolved in the second pass below, once Layer 2's verdicts
            # (if any) are in -- placeholders for now.
            '_duplicate': (duplicate_match, _classify_similarity(duplicate_sim)),
            '_category': (category_match, _classify_similarity(category_sim)),
            '_subcategory': (subcategory_match, _classify_similarity(subcategory_sim)),
            '_supplier': (supplier_match, _classify_similarity(supplier_sim)),
        }
        for kind in ('_duplicate', '_category', '_subcategory', '_supplier'):
            match, status = row[kind]
            if status == 'ambiguous':
                excel_text = {'_duplicate': name, '_category': category_name,
                               '_subcategory': subcategory_name, '_supplier': supplier_name}[kind]
                ai_pending.append({
                    'key': (row_number, kind), 'kind': kind.lstrip('_'),
                    'name': excel_text, 'candidate_name': match.name,
                })
        pending_rows.append(row)

    ai_verdicts = _ai_judge_ambiguous_matches(ai_pending)

    def resolve(row_number, kind, match, status):
        """True (use `match`) if Layer 1 was confident on its own, or Layer
        2 (if it ran for this pair) said yes. Anything else -- no
        candidate, low similarity, ambiguous with no/negative Layer 2
        verdict -- resolves to "don't use it", Layer 1's own conservative
        default."""
        if status == 'confident':
            return True
        if status == 'ambiguous':
            return ai_verdicts.get((row_number, kind), False)
        return False

    rows_out = []
    for row in pending_rows:
        if not row.get('is_new'):
            rows_out.append(row)
            continue

        row_number = row['row_number']
        duplicate_match, duplicate_status = row.pop('_duplicate')
        category_match, category_status = row.pop('_category')
        subcategory_match, subcategory_status = row.pop('_subcategory')
        supplier_match, supplier_status = row.pop('_supplier')

        duplicate = duplicate_match if resolve(row_number, '_duplicate', duplicate_match, duplicate_status) else None
        category = category_match if resolve(row_number, '_category', category_match, category_status) else None
        subcategory = subcategory_match if resolve(row_number, '_subcategory', subcategory_match, subcategory_status) else None
        supplier = supplier_match if resolve(row_number, '_supplier', supplier_match, supplier_status) else None
        resolved_category = subcategory or category

        row['category_id'] = resolved_category.pk if resolved_category else None
        row['category_name'] = resolved_category.name if resolved_category else ''
        row['supplier_id'] = supplier.pk if supplier else None
        row['supplier_name'] = supplier.name if supplier else ''
        row['duplicate_candidate'] = {'id': duplicate.pk, 'name': duplicate.name} if duplicate else None
        rows_out.append(row)

    return {'rows': rows_out, 'skipped': skipped, 'errors': errors}
